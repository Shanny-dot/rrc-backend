"""End-of-day report for staff: today's activity plus every member's points and status, as an
Excel file emailed to ADMIN_EMAIL (run nightly by .github/workflows/daily-report.yml)."""
import datetime
import io
from collections import Counter
from zoneinfo import ZoneInfo

from app.core.supabase_client import supabase_admin

SYDNEY = ZoneInfo("Australia/Sydney")
LAUNCH_CODE = "LAUNCH"


def _local(ts):
    if not ts:
        return None
    return datetime.datetime.fromisoformat(str(ts).replace("Z", "+00:00")).astimezone(SYDNEY)


def _fmt(ts):
    d = _local(ts)
    return d.strftime("%d %b %Y %H:%M") if d else ""


def _all(table: str, cols: str) -> list:
    """Every row of a table (PostgREST returns 1,000 rows at most per request)."""
    rows, start = [], 0
    while True:
        chunk = supabase_admin.table(table).select(cols).range(start, start + 999).execute().data or []
        rows += chunk
        if len(chunk) < 1000:
            return rows
        start += 1000


def member_status(u: dict, tx_sum: dict) -> tuple:
    """(status, plain-English reason) for one member — how they got their points, or why not."""
    verified = bool(u.get("mobile_verified_at") and u.get("email_verified_at"))
    code = (u.get("signup_code") or "").strip().upper()
    if not verified:
        missing = ("mobile and email" if not u.get("mobile_verified_at") and not u.get("email_verified_at")
                   else "mobile number" if not u.get("mobile_verified_at") else "email")
        then = ("Their 100 points are added automatically when they finish." if code == LAUNCH_CODE
                else "No LAUNCH code, so no automatic Welcome Bonus when they finish.")
        return "Not verified yet", f"Hasn't confirmed their {missing}. {then}"
    if tx_sum.get("signup_bonus"):
        return "Bonus - automatic", ("Signed up with the LAUNCH code and verified; Welcome Bonus added automatically."
                                     if code == LAUNCH_CODE else "Welcome Bonus added automatically at sign-up.")
    if tx_sum.get("admin_adjustment", 0) > 0 and code == LAUNCH_CODE:
        return "Bonus - staff", "Had the LAUNCH code; staff added the points by hand."
    if tx_sum.get("admin_adjustment", 0) > 0:
        why = (f"Joined with a friend's code ({code}), not LAUNCH" if code else "Joined without any code")
        return "Bonus - staff", f"{why}; staff added points by hand."
    if code == LAUNCH_CODE and not u.get("welcomed_at"):
        return "Waiting", "Had the LAUNCH code; the bonus is added automatically the next time they open the app."
    if code == LAUNCH_CODE:
        return "Check", "Has the LAUNCH code and is verified but has no Welcome Bonus — please check."
    return "No bonus - no code", "Verified, but joined without the LAUNCH code, so no automatic Welcome Bonus."


def build(day: datetime.date = None) -> tuple:
    """(summary dict, xlsx bytes) for the given Sydney day (default: today in Sydney)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    day = day or datetime.datetime.now(SYDNEY).date()
    is_today = lambda ts: bool(ts) and _local(ts).date() == day  # noqa: E731

    users = [u for u in _all("users", "*") if not u.get("deleted_at")]
    by_id = {u["id"]: u for u in users}
    txs = _all("transactions", "user_id, type, points, description, created_at")
    reds = _all("redemptions", "id, user_id, reward_id, status, points_used, requested_at")
    refs = _all("referrals", "referrer_id, referred_name, referred_email, referred_mobile, status, channel, created_at")
    rewards = {r["id"]: r.get("title") for r in _all("rewards", "id, title")}

    sums = {}
    for t in txs:
        sums.setdefault(t["user_id"], Counter())[t["type"]] += t.get("points") or 0
    refs_made = Counter(r["referrer_id"] for r in refs)

    joined_today = [u for u in users if is_today(u.get("created_at"))]
    tx_today = [t for t in txs if is_today(t.get("created_at"))]
    reds_today = [r for r in reds if is_today(r.get("requested_at"))]
    refs_today = [r for r in refs if is_today(r.get("created_at"))]
    statuses = Counter(member_status(u, sums.get(u["id"], {}))[0] for u in users)
    summary = {
        "date": day.strftime("%A %d %B %Y"),
        "members_total": len(users),
        "members_verified": sum(1 for u in users if u.get("mobile_verified_at") and u.get("email_verified_at")),
        "joined_today": len(joined_today),
        "joined_today_verified": sum(1 for u in joined_today if u.get("mobile_verified_at") and u.get("email_verified_at")),
        "points_added_today": sum(t["points"] for t in tx_today if (t.get("points") or 0) > 0),
        "points_used_today": -sum(t["points"] for t in tx_today if (t.get("points") or 0) < 0),
        "redemptions_today": len(reds_today),
        "redemptions_open": sum(1 for r in reds if r.get("status") in ("pending", "in_progress", "approved")),
        "referrals_today": len(refs_today),
        "statuses": dict(statuses),
    }

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="5B2C83"), Font(bold=True, color="FFFFFF")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Summary"
    labels = [("Report for", summary["date"]), ("Members (total)", summary["members_total"]),
              ("Members fully verified", summary["members_verified"]), ("Joined today", summary["joined_today"]),
              ("Joined today and verified", summary["joined_today_verified"]),
              ("Points added today", summary["points_added_today"]), ("Points used today", summary["points_used_today"]),
              ("Reward requests today", summary["redemptions_today"]),
              ("Reward requests still open (pending / in progress / approved)", summary["redemptions_open"]),
              ("Referrals made today", summary["referrals_today"])]
    labels += [(f"Status: {k}", v) for k, v in sorted(statuses.items())]
    sheet(ws, ["", "Value"], labels, [55, 30])

    yn = lambda b: "Yes" if b else "No"  # noqa: E731
    name = lambda uid: (by_id.get(uid) or {}).get("full_name") or (by_id.get(uid) or {}).get("email") or "(deleted member)"  # noqa: E731
    sheet(wb.create_sheet("Joined today"),
          ["Member", "Email", "Mobile", "Joined", "Signed up with", "Code", "Referred by", "Mobile verified",
           "Email verified", "Points", "Status", "Why"],
          [[u.get("full_name"), u.get("email"), u.get("mobile"), _fmt(u.get("created_at")), u.get("signup_method") or "",
            u.get("signup_code") or "None", name(u["referred_by"]) if u.get("referred_by") else "",
            yn(u.get("mobile_verified_at")), yn(u.get("email_verified_at")), u.get("total_points") or 0,
            *member_status(u, sums.get(u["id"], {}))] for u in sorted(joined_today, key=lambda u: u.get("created_at") or "")],
          [22, 30, 15, 18, 10, 12, 18, 9, 9, 8, 20, 60])
    sheet(wb.create_sheet("Points today"), ["Time", "Member", "Type", "Points", "Description"],
          [[_fmt(t["created_at"]), name(t["user_id"]), t["type"], t["points"], (t.get("description") or "").replace("\n", " ")]
           for t in sorted(tx_today, key=lambda t: t["created_at"])], [18, 24, 18, 9, 60])
    sheet(wb.create_sheet("Reward requests today"), ["Requested", "Member", "Reward", "Points", "Status"],
          [[_fmt(r["requested_at"]), name(r["user_id"]), rewards.get(r["reward_id"], ""), r.get("points_used"), r.get("status")]
           for r in sorted(reds_today, key=lambda r: r["requested_at"])], [18, 24, 32, 9, 14])
    sheet(wb.create_sheet("Referrals today"), ["Made", "Referrer", "Friend", "Friend contact", "How", "Stage"],
          [[_fmt(r["created_at"]), name(r["referrer_id"]), r.get("referred_name"),
            r.get("referred_email") or r.get("referred_mobile") or "", r.get("channel"), r.get("status")]
           for r in sorted(refs_today, key=lambda r: r["created_at"])], [18, 24, 22, 28, 10, 12])

    rows = []
    for u in sorted(users, key=lambda u: u.get("created_at") or ""):
        s = sums.get(u["id"], Counter())
        status, why = member_status(u, s)
        rows.append([u.get("full_name"), u.get("email"), u.get("mobile"), _fmt(u.get("created_at")), u.get("signup_method") or "",
                     u.get("signup_code") or "None", name(u["referred_by"]) if u.get("referred_by") else "",
                     yn(u.get("mobile_verified_at")), yn(u.get("email_verified_at")), yn(u.get("profile_complete")),
                     u.get("total_points") or 0, s.get("signup_bonus", 0), s.get("admin_adjustment", 0),
                     sum(v for k, v in s.items() if "referral" in k), sum(v for k, v in s.items() if "event" in k or "attend" in k),
                     s.get("redemption", 0) + s.get("redemption_refund", 0), refs_made.get(u["id"], 0), status, why])
    sheet(wb.create_sheet("All members"),
          ["Member", "Email", "Mobile", "Joined", "Signed up with", "Code", "Referred by", "Mobile verified", "Email verified",
           "Profile complete", "Points now", "Welcome (auto)", "Added by staff", "Referral bonus", "Events", "Redeemed",
           "Friends referred", "Status", "Why"],
          rows, [22, 30, 15, 18, 10, 12, 18, 9, 9, 9, 9, 9, 9, 9, 8, 9, 9, 18, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return summary, buf.getvalue()
